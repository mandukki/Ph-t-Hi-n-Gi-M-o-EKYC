"""Build the group coursework report from verified repository results."""

from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'report' / 'Bao_cao_nhom_Xu_ly_anh_De_tai_7.docx'
WORKBOOK = Path(r'C:\Users\ManhDuc\Downloads\Phân công bài tập nhóm xử lí ảnh đề tài 7.xlsx')

wb = load_workbook(WORKBOOK, data_only=True, read_only=True)
ws = wb['Thành viên']
members = [(str(int(ws.cell(r, 1).value)), str(ws.cell(r, 2).value),
            str(ws.cell(r, 3).value), str(ws.cell(r, 4).value)) for r in range(4, 9)]

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.5)
sec.left_margin, sec.right_margin = Cm(3), Cm(2)
sec.header_distance, sec.footer_distance = Cm(1.25), Cm(1.25)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Times New Roman'; normal.font.size = Pt(13)
normal.paragraph_format.line_spacing = 1.3
normal.paragraph_format.space_after = Pt(6)
for name, size, before, after in [('Title', 18, 0, 18), ('Heading 1', 16, 18, 10),
                                  ('Heading 2', 14, 12, 7), ('Heading 3', 13, 8, 5)]:
    st = styles[name]
    st.font.name = 'Times New Roman'; st.font.size = Pt(size)
    st.font.bold = True; st.font.color.rgb = None
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

def para(text='', boldlead=None, align=None):
    p = doc.add_paragraph()
    if boldlead and text.startswith(boldlead):
        p.add_run(boldlead).bold = True
        p.add_run(text[len(boldlead):])
    else:
        p.add_run(text)
    if align is not None: p.alignment = align
    return p

def heading(text, level=1): doc.add_heading(text, level=level)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers): t.rows[0].cells[i].text = str(h)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row): cells[i].text = str(val)
    for ri, row in enumerate(t.rows):
        for i, c in enumerate(row.cells):
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths: c.width = Cm(widths[i])
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                for run in p.runs:
                    run.font.name = 'Times New Roman'; run.font.size = Pt(10.5)
                    if ri == 0: run.bold = True
            if ri == 0:
                sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), 'DCE6F1'); c._tc.get_or_add_tcPr().append(sh)
        if ri == 0:
            trPr = row._tr.get_or_add_trPr(); rep=OxmlElement('w:tblHeader'); rep.set(qn('w:val'),'true'); trPr.append(rep)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def figure(path, caption, width=14.5):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ROOT/path), width=Cm(width))
    cap=doc.add_paragraph(caption); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs: run.italic=True; run.font.size=Pt(11)

def page(): doc.add_page_break()

def bullet(text): doc.add_paragraph(text, style='List Bullet')

def field(par, instr):
    b=OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'),'begin')
    x=OxmlElement('w:instrText'); x.set(qn('xml:space'),'preserve'); x.text=instr
    s=OxmlElement('w:fldChar'); s.set(qn('w:fldCharType'),'separate')
    t=OxmlElement('w:t'); t.text='1' if instr == 'PAGE' else 'Nhấn Ctrl+A, F9 để cập nhật mục lục.'
    e=OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'),'end')
    for el in (b,x,s,t,e): par.add_run()._r.append(el)

# Cover
para('ĐẠI HỌC KINH TẾ QUỐC DÂN', align=WD_ALIGN_PARAGRAPH.CENTER)
para('TRƯỜNG CÔNG NGHỆ', align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph('\n\n\n')
p=para('BÁO CÁO BÀI TẬP NHÓM', align=WD_ALIGN_PARAGRAPH.CENTER)
for r in p.runs: r.bold=True; r.font.size=Pt(19)
p=para('MÔN XỬ LÝ ẢNH', align=WD_ALIGN_PARAGRAPH.CENTER)
for r in p.runs: r.bold=True; r.font.size=Pt(17)
doc.add_paragraph('\n')
p=para('ĐỀ TÀI 7', align=WD_ALIGN_PARAGRAPH.CENTER)
for r in p.runs: r.bold=True; r.font.size=Pt(16)
p=para('PHÁT HIỆN GIAN LẬN VÀ XÁC THỰC GIẤY TỜ CÁ NHÂN (eKYC) BẰNG PHÂN TÍCH KẾT CẤU BỀ MẶT VÀ GỢI Ý VÙNG CAN THIỆP ẢNH', align=WD_ALIGN_PARAGRAPH.CENTER)
for r in p.runs: r.bold=True; r.font.size=Pt(16)
doc.add_paragraph('\n\n')
para('Giảng viên hướng dẫn: Phạm Minh Hoàn')
para('Nhóm thực hiện: Phạm Mạnh Đức, Bạch Quốc Khánh, Trịnh Mai Linh, Trần Đức Mạnh, Vũ Bình Minh')
doc.add_paragraph('\n\n')
para('Hà Nội, 2026', align=WD_ALIGN_PARAGRAPH.CENTER)
page()

heading('THÔNG TIN NHÓM VÀ PHÂN CÔNG',1)
para('Vai trò dưới đây lấy từ bảng phân công của nhóm. Mức độ đóng góp thực tế cần các thành viên xác nhận trước khi nộp; bảng kế hoạch không chứng minh tỷ lệ đóng góp cá nhân.')
table(['STT','Mã sinh viên','Họ tên','Vai trò / phần việc dự kiến'],
      [(i,code,name,detail) for i,(code,name,role,detail) in enumerate(members,1)], [1,2.6,3.7,9.5])
heading('MỤC LỤC',1)
para('Trong Microsoft Word, chọn toàn bộ văn bản bằng Ctrl+A rồi nhấn F9 để cập nhật mục lục và số trang.')
field(doc.add_paragraph(),'TOC \\o "1-3" \\h \\z \\u')
page()

heading('DANH MỤC TỪ VIẾT TẮT',1)
table(['Từ viết tắt','Ý nghĩa'], [('eKYC','Electronic Know Your Customer, định danh khách hàng điện tử'),('LBP','Local Binary Patterns'),('GLCM','Gray-Level Co-occurrence Matrix'),('FFT','Fast Fourier Transform'),('SVM','Support Vector Machine'),('RF','Random Forest'),('LOGO','Leave-One-Group-Out'),('AUC','Area Under the ROC Curve'),('EER','Equal Error Rate'),('FAR','Tỷ lệ thẻ giả bị chấp nhận'),('FRR','Tỷ lệ thẻ thật bị từ chối')],[3.2,13.6])
heading('LỜI MỞ ĐẦU',1)
para('Trong quy trình định danh điện tử, ảnh giấy tờ có thể bị sửa chữ, thay vùng chân dung, cắt ghép hoặc in và chụp lại. Các thao tác này có thể làm thay đổi cấu trúc cục bộ, quan hệ mức xám và phân bố tần số của ảnh. Nhóm xây dựng một hệ thống thử nghiệm để phân loại ảnh giấy tờ thật/giả bằng đặc trưng kết cấu, đồng thời gợi ý những ô ảnh cần kiểm tra kỹ hơn.')
para('Báo cáo mô tả bộ dữ liệu 165 ảnh, quy trình tiền xử lý, đặc trưng LBP–GLCM–FFT, mô hình SVM và Random Forest, cách đánh giá theo nhóm định danh và kết quả trên tập dữ liệu hiện tại. Kết quả được trình bày cùng các hạn chế để người đọc phân biệt khả năng thử nghiệm với khả năng ứng dụng thực tế.')
page()

heading('CHƯƠNG 1. GIỚI THIỆU ĐỀ TÀI',1)
heading('1.1. Bối cảnh và bài toán',2)
para('Đầu vào của hệ thống là ảnh mặt trước giấy tờ cá nhân. Đầu ra chính là dự đoán REAL hoặc FAKE; đầu ra hỗ trợ là bản đồ điểm nghi vấn theo lưới 4 × 6. Bài toán tập trung vào dấu vết thị giác trong ảnh, không kiểm tra pháp lý của giấy tờ, tính hợp lệ của số định danh hoặc đối chiếu khuôn mặt với người sử dụng.')
heading('1.2. Mục tiêu và phạm vi',2)
bullet('Xây dựng tập ảnh gồm ảnh thật và bốn kiểu giả mạo; lưu nhãn cùng nhóm nguồn để đánh giá tránh rò rỉ.')
bullet('Chuẩn hóa ảnh 856 × 540, chuyển mức xám và trích xuất 60 đặc trưng LBP, GLCM, FFT.')
bullet('So sánh mô hình SVM, Random Forest với đối chứng LBP-only và lớp đa số.')
bullet('Đánh giá theo 5 nhóm nguồn, phân tích riêng các kiểu can thiệp và thăm dò gợi ý ô ảnh bị sửa.')
heading('1.3. Sản phẩm của nhóm',2)
para('Repo gồm mã nguồn tiền xử lý, sinh dữ liệu giả số hóa, trích xuất đặc trưng, huấn luyện, đánh giá, demo một ảnh, bộ kiểm thử và báo cáo kết quả. Script demo lưu ảnh overlay và JSON dự đoán; đây là bản thử nghiệm dòng lệnh, chưa phải ứng dụng eKYC vận hành trong thực tế.')

heading('CHƯƠNG 2. CƠ SỞ KỸ THUẬT',1)
heading('2.1. Kết cấu ảnh và tiền xử lý',2)
para('Ảnh chụp chịu ảnh hưởng của phối cảnh, ánh sáng, độ nét và thiết bị. Module tiền xử lý tìm vùng thẻ bằng contour, áp dụng phép biến đổi phối cảnh nếu tìm được tứ giác phù hợp, sau đó chuẩn hóa kích thước và chuyển ảnh xám. Cách chuẩn hóa này giúp các bước trích xuất đặc trưng làm việc trên cùng hệ tọa độ; khi cắt tự động thất bại, hành vi fallback của code cần được lưu ý khi diễn giải kết quả.')
heading('2.2. LBP, GLCM và FFT',2)
para('LBP so sánh mức xám của các điểm lân cận với tâm và tổng hợp mã nhị phân thành histogram 10 chiều. GLCM đếm quan hệ đồng xuất hiện của các mức xám theo khoảng cách và hướng; pipeline lấy 6 thuộc tính ở 3 khoảng cách, cho 18 chiều. FFT chuyển ảnh sang miền tần số và tóm tắt năng lượng theo 32 vòng bán kính. Ba nhóm được nối thành vector 60 chiều. Cách thực thi cụ thể nằm tại src/features, vì thế các khái niệm ở đây được trình bày ở mức đủ để giải thích pipeline [1].')
heading('2.3. Bộ phân loại và chỉ số đánh giá',2)
para('SVM tìm ranh giới phân lớp trong không gian đặc trưng; Random Forest tổng hợp nhiều cây quyết định. Accuracy là tỷ lệ dự đoán đúng; Balanced Accuracy lấy trung bình recall theo hai lớp để giảm ảnh hưởng mất cân bằng. Precision, recall và F1 trong báo cáo lấy FAKE làm lớp dương. AUC-ROC đo năng lực xếp hạng theo điểm số khi thay đổi ngưỡng. EER là mức lỗi tại điểm hai loại lỗi xấp xỉ nhau. FAR ở repo là tỷ lệ thẻ giả lọt qua, còn FRR là tỷ lệ thẻ thật bị chặn nhầm [2].')

heading('CHƯƠNG 3. DỮ LIỆU VÀ TIỀN XỬ LÝ',1)
heading('3.1. Nguồn và cơ cấu dữ liệu',2)
para('Bộ dữ liệu hiện có 165 ảnh: 45 ảnh thật và 120 ảnh giả. Theo xác nhận của nhóm, chủ thể của ảnh thẻ thật đã đồng ý sử dụng cho bài tập; repo không lưu văn bản chứng minh sự đồng ý. Có 5 nhóm thẻ nguồn. 90 ảnh giả số hóa chia đều cho edit_text, replace_photo và splice; 30 ảnh reprint được thu thập riêng và đã gán group thủ công theo thông tin hiển thị trên thẻ. Đây là dữ liệu nội bộ của bài tập, không phải tập benchmark công khai.')
table(['Loại ảnh','Số mẫu','Mô tả'], [('real / none',45,'Ảnh chụp thẻ thật'),('edit_text',30,'Sửa vùng chữ hoặc số'),('replace_photo',30,'Can thiệp vùng ảnh chân dung'),('splice',30,'Cắt ghép vùng nội dung'),('reprint',30,'In/hiển thị lại rồi chụp')],[3.2,2.2,11.4])
para('Bảng 3.1. Phân bố theo loại ảnh (nguồn: data/labels.csv và evaluation_summary.md).',align=WD_ALIGN_PARAGRAPH.CENTER)
table(['Nhóm nguồn','Tổng','Thật','Giả','Reprint'], [('group_01',34,9,25,7),('group_02',27,7,20,2),('group_03',38,11,27,9),('group_04',31,10,21,3),('group_05',35,8,27,9)],[4.1,2,2,2,2.5])
para('Bảng 3.2. Phân bố theo nhóm nguồn. Reprint chưa cân bằng giữa các nhóm.',align=WD_ALIGN_PARAGRAPH.CENTER)
heading('3.2. Quy trình xử lý',2)
para('Mỗi bản ghi có path, label, fraud_type và group_id. Pipeline đọc ảnh, chuẩn hóa mặt thẻ 856 × 540 và chuyển grayscale trước khi trích xuất đặc trưng. Chế độ --strict dừng khi gặp ảnh lỗi. Bộ đánh giá kiểm tra đường dẫn, nhãn, nhóm nguồn và phát hiện ảnh trùng byte bằng SHA-256 trước khi huấn luyện. Tập đặc trưng đã được xây dựng thành công cho 165/165 ảnh.')
heading('3.3. Bảo vệ dữ liệu',2)
para('Báo cáo này không chèn ảnh giấy tờ thật hoặc thông tin định danh. Khi trình chiếu demo, nhóm cần dùng ảnh được cho phép chia sẻ và che dữ liệu cá nhân. Repo chặn ảnh dữ liệu bằng .gitignore; các file ảnh trong bản ZIP trao đổi nội bộ vẫn cần được kiểm soát người nhận.')

heading('CHƯƠNG 4. THIẾT KẾ VÀ XÂY DỰNG HỆ THỐNG',1)
heading('4.1. Kiến trúc xử lý',2)
para('Luồng xử lý: ảnh đầu vào → tìm/cắt thẻ → ảnh xám 856 × 540 → LBP (10) + GLCM (18) + FFT (32) → vector 60 chiều → SVM/RF → nhãn thật/giả. Nhánh gợi ý vùng can thiệp chia ảnh thành 24 ô, trích đặc trưng từng ô và xếp hạng các ô có dấu hiệu sửa.')
table(['Thành phần','File/module','Chức năng'], [('Dữ liệu','data_loader.py; labels.csv','Đọc ảnh, nhãn, nhóm nguồn'),('Tiền xử lý','src/preprocess/perspective.py','Cắt và chuẩn hóa thẻ'),('Đặc trưng','src/features/*.py','LBP, GLCM, FFT, ghép vector'),('Mô hình','src/models/*.py','SVM, RF, ngưỡng và chỉ số'),('Đánh giá','scripts/run_grouped_evaluation.py','LOGO, báo cáo và biểu đồ'),('Định vị','src/localization.py','Điểm nghi vấn 24 ô'),('Demo','scripts/demo_image.py','Dự đoán một ảnh và xuất overlay')],[3,5.5,8.3])
para('Bảng 4.1. Các thành phần chính của repo.',align=WD_ALIGN_PARAGRAPH.CENTER)
heading('4.2. Thiết kế đánh giá chống rò rỉ',2)
para('Mỗi fold giữ lại toàn bộ ảnh của một group để kiểm tra và huấn luyện trên bốn group còn lại. Bản thật và các ảnh giả sinh từ cùng thẻ không thể đồng thời nằm ở train và test. Trong tập train của mỗi fold ngoài, ngưỡng phân loại được ước lượng từ dự đoán ngoài fold theo nhóm. Tập test ngoài không tham gia học ngưỡng. Với SVM, điểm dùng để đánh giá là decision_function, không được diễn giải là xác suất đã hiệu chuẩn [2].')
heading('4.3. Gợi ý vùng can thiệp',2)
para('Nhánh định vị chia ảnh chuẩn hóa thành lưới 4 × 6. Random Forest theo ô được học từ các vùng sửa xấp xỉ do script tạo ảnh giả cung cấp. Hệ thống xếp hạng ô và vẽ overlay. Tọa độ là nhãn yếu, chưa được con người xác minh theo pixel và có thể lệch sau phép biến đổi phối cảnh. Vì vậy output là vùng cần kiểm tra, không phải mask chính xác.')

heading('CHƯƠNG 5. THỰC NGHIỆM VÀ KẾT QUẢ',1)
heading('5.1. Thiết lập',2)
para('Các thí nghiệm dùng 165 mẫu và 5 fold Leave-One-Group-Out. Nhóm so sánh lớp đa số, LBP-only 10 chiều và vector đầy đủ 60 chiều, mỗi bộ đặc trưng với SVM hoặc Random Forest. Bảng macro là trung bình không trọng số theo 5 group; các chỉ số pooled gộp dự đoán ngoài fold của 165 ảnh. Hai cách tổng hợp có thể cho số khác nhau.')
heading('5.2. So sánh đối chứng đặc trưng',2)
table(['Đặc trưng','Mô hình','Accuracy macro','Balanced Acc','AUC macro','EER macro'], [('LBP 10','SVM','50,6%','52,2%','0,588','0,432'),('LBP 10','RF','53,9%','53,5%','0,564','0,437'),('Đầy đủ 60','SVM','65,5%','60,6%','0,693','0,311'),('Đầy đủ 60','RF','65,8%','68,2%','0,757','0,309')],[3.1,1.7,3.1,3.1,2.8,2.8])
para('Bảng 5.1. Đối chứng LBP-only với LBP + GLCM + FFT trên cùng 5 fold.',align=WD_ALIGN_PARAGRAPH.CENTER)
para('Vector đầy đủ cải thiện các chỉ số so với LBP-only trên cả hai mô hình. Tuy nhiên accuracy macro 65,5–65,8% vẫn thấp hơn đối chứng luôn dự đoán FAKE là 72,7%. Do đó không thể dùng accuracy riêng lẻ để khẳng định hệ thống đã phân loại tốt. RF đạt Balanced Accuracy macro 68,2% và AUC macro 0,757; SVM đạt lần lượt 60,6% và 0,693.')
heading('5.3. Kết quả phân loại theo lớp và loại can thiệp',2)
table(['Mô hình','Precision fake','Recall fake','Specificity real','F1 fake','FAR','FRR'], [('SVM','79,6%','71,7%','51,1%','75,4%','28,3%','48,9%'),('RF','85,7%','65,0%','71,1%','73,9%','35,0%','28,9%')],[2,2.5,2.5,2.7,2.2,2.1,2.1])
para('Bảng 5.2. Chỉ số pooled ngoài fold trên 165 mẫu.',align=WD_ALIGN_PARAGRAPH.CENTER)
table(['Loại','SVM đúng/tổng','RF đúng/tổng'], [('edit_text','23/30 (76,7%)','19/30 (63,3%)'),('replace_photo','21/30 (70,0%)','20/30 (66,7%)'),('splice','22/30 (73,3%)','20/30 (66,7%)'),('reprint','20/30 (66,7%)','19/30 (63,3%)'),('real / none','23/45 (51,1%)','32/45 (71,1%)')],[5.5,5.4,5.4])
para('Bảng 5.3. Tỷ lệ phát hiện từng loại; hàng real thể hiện tỷ lệ nhận đúng ảnh thật.',align=WD_ALIGN_PARAGRAPH.CENTER)
figure(Path('report/evaluation/svm_confusion_matrix.png'),'Hình 5.1. Ma trận nhầm lẫn ngoài fold của SVM.',12.8)
figure(Path('report/evaluation/rf_confusion_matrix.png'),'Hình 5.2. Ma trận nhầm lẫn ngoài fold của Random Forest.',12.8)
heading('5.4. Ảnh hưởng của reprint',2)
table(['Mô hình','Tập','Số ảnh','Accuracy macro','Balanced Acc','Specificity real'], [('SVM','Full',165,'65,5%','60,6%','51,1%'),('SVM','Digital-only',135,'78,0%','78,7%','77,8%'),('RF','Full',165,'65,8%','68,2%','71,1%'),('RF','Digital-only',135,'71,8%','73,7%','77,8%')],[2.2,3.1,2,3.5,3.1,3.3])
para('Bảng 5.4. Đánh giá độ nhạy khi bỏ 30 ảnh reprint. Hai tập có cơ cấu và bài toán khác nhau, không phải phép thử cùng một mô hình trên cùng test.',align=WD_ALIGN_PARAGRAPH.CENTER)
para('Trên tập chỉ gồm ảnh thật và ảnh giả số hóa, SVM đạt accuracy macro 78,0%. Chênh lệch với Full gợi ý reprint là nguồn khó khăn, nhưng việc bỏ 30 ảnh cũng thay đổi kích thước và thành phần dữ liệu. Phân bố reprint giữa các group là 7–2–9–3–9 nên cần thêm dữ liệu cân bằng trước khi đưa ra kết luận rộng hơn.')
heading('5.5. Gợi ý vùng can thiệp',2)
table(['Chỉ số','Bộ phân loại ô','Chọn ô ngẫu nhiên'], [('Top-1 hit','70,0%','19,4%'),('Top-3 hit','84,4%','48,2%')],[5.3,5.5,5.5])
para('Bảng 5.5. Đánh giá trên 90 ảnh giả số hóa với nhãn vùng xấp xỉ từ script sinh ảnh.',align=WD_ALIGN_PARAGRAPH.CENTER)
para('Mô hình phát hiện ô chứa vùng sửa ở vị trí Top-1 trong 70,0% và Top-3 trong 84,4% trường hợp. Các mức chọn ngẫu nhiên tương ứng là 19,4% và 48,2%. Tỷ lệ Top-1 theo loại là edit_text 86,7%, replace_photo 66,7% và splice 56,7%. Không có đánh giá cho reprint vì chưa có mask; kết quả chưa phải độ chính xác phân đoạn ở mức pixel.')

heading('CHƯƠNG 6. DEMO VÀ HƯỚNG DẪN CHẠY',1)
heading('6.1. Chuẩn bị môi trường',2)
para('Trên Windows, mở thư mục dự án bằng Visual Studio Code rồi dùng Terminal PowerShell. Cài Python, tạo môi trường ảo và cài các gói trong requirements.txt. Cần đặt dữ liệu ảnh và labels.csv đúng đường dẫn mô tả ở README vì ảnh nhạy cảm không nằm trong Git.')
for line in ['python -m venv .venv', '.venv\\Scripts\\Activate.ps1', 'pip install -r requirements.txt', 'python -m unittest discover -s tests -p "test_*.py"']:
    p=para(line); p.paragraph_format.left_indent=Cm(0.7)
heading('6.2. Tái tạo kết quả',2)
for line in ['python scripts/build_features.py --strict', 'python scripts/run_grouped_evaluation.py --model all', 'python scripts/run_lbp_baseline.py', 'python scripts/evaluate_patch_classifier.py', 'python scripts/demo_image.py --image "duong_dan_anh_moi.jpg"']:
    p=para(line); p.paragraph_format.left_indent=Cm(0.7)
para('Nhánh demo in nhãn, tạo JSON và ảnh overlay trong data/processed. Nên dùng ảnh mới không có trong labels.csv khi trình diễn. Script tổ chức reprint chỉ dành cho bước dựng lại dữ liệu từ đầu; bộ hiện tại đã xác minh group thủ công nên không chạy lại bước gán heuristic.')
heading('6.3. Kiểm thử và lưu vết',2)
para('Ở lần kiểm tra gần nhất, 42/42 bài kiểm thử tự động đã đạt; xây dựng đặc trưng thành công cho 165/165 mẫu. Evaluation manifest lưu hash dữ liệu, mã nguồn, thông số và trạng thái Git. Vì lần chạy hiện tại ở worktree chưa commit, commit SHA một mình không đại diện đầy đủ cho mã đã chạy; cần giữ manifest và các file nguồn tương ứng khi tái lập.')

heading('CHƯƠNG 7. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN',1)
heading('7.1. Kết quả đạt được',2)
para('Nhóm đã hoàn thiện pipeline đặc trưng kết cấu 60 chiều, hai bộ phân loại, đánh giá theo nhóm định danh, đối chứng LBP-only, phân tích theo loại gian lận và demo gợi ý ô nghi vấn. Vector đầy đủ cho kết quả tốt hơn LBP-only trên bộ dữ liệu hiện có. Random Forest cho AUC macro 0,757, trong khi SVM cho 0,693.')
heading('7.2. Hạn chế',2)
para('Bộ dữ liệu chỉ có 5 nhóm thẻ nguồn, 165 ảnh và ảnh giả phần lớn do nhóm tự tạo. Accuracy Full thấp hơn lớp đa số; tỷ lệ reprint theo nhóm không đều. Replace_photo dùng nguồn trong cùng nhóm, chưa mô phỏng đầy đủ thay chân dung bằng danh tính độc lập. Nhãn vùng là xấp xỉ từ chương trình sinh ảnh, chưa có mask xác minh thủ công; chưa thể báo cáo IoU/Dice hoặc hiệu năng trên dữ liệu ngoài tập mô phỏng. Dữ liệu đồng ý sử dụng được nhóm xác nhận nhưng chưa có hồ sơ đồng ý trong repo.')
heading('7.3. Hướng phát triển',2)
para('Mở rộng số nhóm thẻ, thiết bị chụp và điều kiện ánh sáng; cân bằng reprint; chuẩn bị tập kiểm định độc lập; gán mask vùng can thiệp bởi người kiểm tra; thử các phép cắt thẻ ổn định hơn và so sánh thêm phương pháp học sâu gọn nhẹ. Khi có dữ liệu mới, giữ nguyên nguyên tắc chia train/test theo nhóm nguồn.')

heading('TÀI LIỆU THAM KHẢO',1)
para('[1] scikit-image, “skimage.feature API: local_binary_pattern, graycomatrix, graycoprops,” https://scikit-image.org/docs/stable/api/skimage.feature.html (truy cập 29/10/2026).')
para('[2] scikit-learn, “Cross-validation: evaluating estimator performance” và “Metrics and scoring,” https://scikit-learn.org/stable/modules/cross_validation.html; https://scikit-learn.org/stable/modules/model_evaluation.html (truy cập 29/10/2026).')
para('[3] Nhóm đề tài 7, “README và mã nguồn dự án eKYC Texture Fraud Detection,” bản làm việc ngày 28/10/2026, lưu tại repo nội bộ.')
para('[4] Nhóm đề tài 7, “evaluation_summary.md, lbp_baseline/summary.md, patch_classifier/summary.md,” kết quả thực nghiệm trong report/evaluation, 28/10/2026.')
heading('PHỤ LỤC. DANH SÁCH OUTPUT',1)
table(['Output','Vị trí trong repo'], [('Bảng tổng hợp đánh giá','report/evaluation/evaluation_summary.md'),('Đối chứng LBP','report/evaluation/lbp_baseline/summary.md'),('Gợi ý ô can thiệp','report/evaluation/patch_classifier/summary.md'),('Ma trận nhầm lẫn và biểu đồ','report/evaluation/*.png'),('Lưu vết đánh giá','report/evaluation/evaluation_manifest.json'),('Demo một ảnh','scripts/demo_image.py')],[6,10.8])

footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field(footer,'PAGE')
settings = doc.settings.element
up=OxmlElement('w:updateFields'); up.set(qn('w:val'),'true'); settings.append(up)
doc.save(OUT)
print(OUT)
